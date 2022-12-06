# Copyright 2022 Abhishek Harshvardhan Mishra
#
# This program is free software: you can redistribute it and/or modify
# it under the terms of the GNU General Public License as published by
# the Free Software Foundation, either version 3 of the License, or
# (at your option) any later version.
#
# This program is distributed in the hope that it will be useful,
# but WITHOUT ANY WARRANTY; without even the implied warranty of
# MERCHANTABILITY or FITNESS FOR A PARTICULAR PURPOSE.  See the
# GNU General Public License for more details.
#
# You should have received a copy of the GNU General Public License
# along with this program.  If not, see <http://www.gnu.org/licenses/>.
import os
import re
import argparse
import sys
import io
import datetime

from bs4 import BeautifulSoup
from docx import Document
from openpyxl import load_workbook
from pdfminer.converter import TextConverter
from pdfminer.layout import LAParams
from pdfminer.pdfdocument import PDFDocument
from pdfminer.pdfinterp import PDFResourceManager, PDFPageInterpreter
from pdfminer.pdfpage import PDFPage


def scrape_html(filepath: str) -> str:
    """Scrape text from an HTML file."""
    with open(filepath, encoding="utf-8") as f:
        soup = BeautifulSoup(f, "html.parser")
        text = soup.get_text(separator="\n")
        return text


def scrape_docx(filepath: str) -> str:
    """Scrape text from a DOCX file."""
    document = Document(filepath)
    text = []
    for p in document.paragraphs:
        text.append(p.text)
    return "\n".join(text)


def scrape_xlsx(filepath: str) -> str:
    """Scrape text from an XLSX file."""
    wb = load_workbook(filepath)
    sheet = wb.active
    text = []
    for row in sheet.rows:
        row_text = []
        for cell in row:
            # Convert floating-point numbers, integers, and datetime
            # values to strings before appending them to the row_text list.
            if isinstance(cell.value, (float, int, datetime.datetime)):
                row_text.append(str(cell.value))
            else:
                row_text.append(cell.value)
        text.append("\t".join(row_text))
    return "\n".join(text)


def scrape_txt(filepath: str) -> str:
    """Scrape text from a TXT file."""
    with open(filepath) as f:
        text = f.read()
        return text


def scrape_csv(filepath: str) -> str:
    """Scrape text from a CSV file."""
    with open(filepath) as f:
        lines = f.readlines()
        text = []
        for line in lines:
            text.append(line.strip())
        return "\n".join(text)


def scrape_pdf(filepath: str) -> str:
    """Scrape text from a PDF file."""
    with open(filepath, "rb") as f:
        resource_manager = PDFResourceManager()

        # Use a StringIO object as the output file object.
        outfp = io.StringIO()
        converter = TextConverter(
            resource_manager, outfp, laparams=LAParams()
        )

        page_interpreter = PDFPageInterpreter(
            resource_manager, converter
        )

        text = []
        for page in PDFPage.get_pages(f):
            page_interpreter.process_page(page)
            text.append(outfp.getvalue())

        return "\n".join(text)


def scrape_file(filepath: str) -> str:
    """Scrape text from a file at the specified filepath.
    Return the scraped text as a string.
    """
    _, ext = os.path.splitext(filepath)
    ext = ext.lower()
    if ext in [".html", ".htm"]:
        return scrape_html(filepath)
    elif ext == ".docx":
        return scrape_docx(filepath)
    elif ext == ".xlsx":
        return scrape_xlsx(filepath)
    elif ext == ".txt":
        return scrape_txt(filepath)
    elif ext == ".csv":
        return scrape_csv(filepath)
    elif ext == ".pdf":
        return scrape_pdf(filepath)
    else:
        raise ValueError(f"Unsupported file type: {ext}")


def main() -> None:
    # Parse the command-line arguments.
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "filepath",
        help="The file to scrape text from.",
    )
    args = parser.parse_args()

    # Scrape the text from the specified file.
    text = scrape_file(args.filepath)

    # Print the scraped text to the console.
    sys.stdout.buffer.write(text.encode("utf-8"))


if __name__ == "__main__":
    main()